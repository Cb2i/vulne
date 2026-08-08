"""Builds Excel workbooks from findings / KPI data for offline reporting, using openpyxl."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.asset import Asset
from app.models.finding import Finding
from app.models.team import Team
from app.schemas.kpi import KPISummary

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

FINDING_COLUMNS = [
    ("Asset", "hostname"),
    ("Plugin", "plugin_name"),
    ("CVE", "cve"),
    ("CVSS", "cvss"),
    ("Score CAA", "caa_score"),
    ("Sévérité CAA", "caa_severity"),
    ("Délai SLE (jours)", "sle_delay_days"),
    ("Échéance SLE", "sle_due_date"),
    ("Équipe", "team_name"),
    ("Statut", "status"),
    ("Ticket", "ticket_ref"),
    ("Première détection", "first_seen_at"),
    ("Dernière détection", "last_seen_at"),
    ("Solution", "solution"),
]


def _autosize(ws) -> None:
    for i, _ in enumerate(ws[1], start=1):
        col_letter = get_column_letter(i)
        max_len = max((len(str(c.value)) for c in ws[col_letter] if c.value is not None), default=10)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)


def export_findings_to_excel(db: Session, findings: list[Finding]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vulnérabilités"

    asset_ids = {f.asset_id for f in findings}
    team_ids = {f.team_id for f in findings if f.team_id}
    assets = {a.id: a for a in db.query(Asset).filter(Asset.id.in_(asset_ids))} if asset_ids else {}
    teams = {t.id: t for t in db.query(Team).filter(Team.id.in_(team_ids))} if team_ids else {}

    headers = [h for h, _ in FINDING_COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for f in findings:
        asset = assets.get(f.asset_id)
        team = teams.get(f.team_id) if f.team_id else None
        row = {
            "hostname": asset.hostname if asset else "",
            "plugin_name": f.plugin_name,
            "cve": f.cve or "",
            "cvss": f.cvss,
            "caa_score": f.caa_score,
            "caa_severity": f.caa_severity.value if f.caa_severity else "",
            "sle_delay_days": f.sle_delay_days,
            "sle_due_date": f.sle_due_date.isoformat() if f.sle_due_date else "",
            "team_name": team.name if team else "Non assigné",
            "status": f.status.value,
            "ticket_ref": f.ticket_ref or "",
            "first_seen_at": f.first_seen_at.strftime("%Y-%m-%d") if f.first_seen_at else "",
            "last_seen_at": f.last_seen_at.strftime("%Y-%m-%d") if f.last_seen_at else "",
            "solution": (f.solution or "")[:500],
        }
        ws.append([row[key] for _, key in FINDING_COLUMNS])

    _autosize(ws)
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_kpi_to_excel(summary: KPISummary) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"

    ws.append(["Indicateur", "Valeur"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    rows = [
        ("Findings ouverts", summary.total_open_findings),
        ("Actifs totaux", summary.total_assets),
        ("Actifs décommissionnés", summary.total_decommissioned_assets),
        ("Exceptions actives", summary.active_exceptions),
        ("Findings en retard (SLE)", summary.overdue_findings),
        ("Taux de conformité SLE (%)", summary.sle_compliance_rate),
        ("Remédiés (30 derniers jours)", summary.remediated_last_30_days),
        ("Nouveaux (30 derniers jours)", summary.new_last_30_days),
    ]
    for label, value in rows:
        ws.append([label, value])

    ws2 = wb.create_sheet("Par sévérité")
    ws2.append(["Sévérité CAA", "Nombre"])
    for cell in ws2[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for sc in summary.by_severity:
        ws2.append([sc.severity, sc.count])

    ws3 = wb.create_sheet("Par équipe")
    ws3.append(["Équipe", "Nombre", "En retard"])
    for cell in ws3[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for tc in summary.by_team:
        ws3.append([tc.team_name, tc.count, tc.overdue_count])

    for sheet in (ws, ws2, ws3):
        _autosize(sheet)
        sheet.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
