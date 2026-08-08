"""Regression tests for importing the 'Exceptions' and 'KPI' sheets from a Tenable
workbook: they must be persisted as reusable ExceptionRecord / OwnershipRule rows that
apply immediately to the findings in the same import, and keep applying to future ones."""
from io import BytesIO

from openpyxl import Workbook

from tests.test_import import _build_workbook


def _build_workbook_with_rules(
    scan_name: str = "VM-TEST",
    hostname: str = "q212675",
    plugin_name: str = "OpenSSL Vulnerability",
) -> bytes:
    base = _build_workbook(scan_name=scan_name, hostname=hostname, plugin_name=plugin_name)
    wb = Workbook()
    wb.remove(wb.active)

    # Reload the base workbook's sheets by re-creating them (openpyxl can't easily merge
    # two in-memory workbooks), so just rebuild the Vulnérabilités + tags sheets here too.
    import openpyxl

    src = openpyxl.load_workbook(BytesIO(base))
    for name in src.sheetnames:
        src_ws = src[name]
        dst_ws = wb.create_sheet(name)
        for row in src_ws.iter_rows(values_only=True):
            dst_ws.append(row)

    exc_ws = wb.create_sheet("Exceptions")
    exc_ws.append([None, None, None, None, None, None])
    exc_ws.append([None, "EXCEPTIONS", None, None, None, None])
    exc_ws.append([None, "Vulnérabilité", "Type d'actif", "Hostname", "Date et porteur", "Lien confluence"])
    exc_ws.append(
        [
            None,
            plugin_name,
            "Poste de travail",
            hostname,
            "27 Janvier 2023 par Eric Mercier",
            "https://example.atlassian.net/wiki/pages/123",
        ]
    )

    kpi_ws = wb.create_sheet("KPI")
    kpi_ws.append([None, None, None, None])
    kpi_ws.append([None, None, None, None])
    kpi_ws.append([None, "Equipes", "Champ_nom d'hotes", "Champ_OS"])
    kpi_ws.append([None, "Postes de travail", "commence par q2", None])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parser_extracts_exceptions_and_ownership_rules():
    from app.importers.tenable_excel.parser import parse_tenable_workbook

    parsed = parse_tenable_workbook(_build_workbook_with_rules(), "test.xlsx")

    assert len(parsed.exceptions) == 1
    exc = parsed.exceptions[0]
    assert exc.vulnerability_match == "OpenSSL Vulnerability"
    assert exc.hostname_pattern == "q212675"
    assert exc.granted_by == "Eric Mercier"
    assert exc.granted_at.isoformat() == "2023-01-27"

    assert len(parsed.ownership_rules) == 1
    rule = parsed.ownership_rules[0]
    assert rule.team_name == "Postes de travail"
    assert rule.hostname_pattern == "commence par q2"


def test_exceptions_sheet_with_trailing_mini_tables_is_not_misparsed():
    """Regression test for a real-world workbook shape: the 'Exceptions' sheet contained
    a second 'Utilitaires/État/Commentaires' blanket-exception table and a third
    wildcard-decommission table below the real one, separated by blank rows. Only the
    first, properly-shaped table must be imported; the rest must surface as a warning
    instead of being silently absorbed as bogus exception rows."""
    from app.importers.tenable_excel.parser import parse_tenable_workbook

    base = _build_workbook(hostname="host01", plugin_name="Real Vulnerability")
    import openpyxl

    src = openpyxl.load_workbook(BytesIO(base))
    wb = Workbook()
    wb.remove(wb.active)
    for name in src.sheetnames:
        dst_ws = wb.create_sheet(name)
        for row in src[name].iter_rows(values_only=True):
            dst_ws.append(row)

    exc_ws = wb.create_sheet("Exceptions")
    exc_ws.append([None] * 6)
    exc_ws.append([None, "EXCEPTIONS", None, None, None, None])
    exc_ws.append([None, "Vulnérabilité", "Type d'actif", "Hostname", "Date et porteur", "Lien confluence"])
    exc_ws.append([None, "Real Vulnerability", "Serveur", "host01", "27 Janvier 2023 par Eric Mercier", "https://x/1"])
    exc_ws.append([None] * 6)  # blank separator -> table ends here
    exc_ws.append([None] * 6)
    exc_ws.append([None, "Utilitaires", "État", "Commentaires", None, None])
    exc_ws.append([None, "contient Security Updates", "Mises a jour automatique", "En fonction de la sortie", None, None])

    buffer = BytesIO()
    wb.save(buffer)

    parsed = parse_tenable_workbook(buffer.getvalue(), "test.xlsx")

    assert len(parsed.exceptions) == 1
    assert parsed.exceptions[0].vulnerability_match == "Real Vulnerability"
    assert any("contenu supplémentaire" in w for w in parsed.warnings)


def test_import_applies_exceptions_and_ownership_rules_immediately(client, admin_token):
    files = {
        "file": (
            "scan.xlsx",
            _build_workbook_with_rules(hostname="q212675", plugin_name="OpenSSL Vulnerability"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    resp = client.post("/api/scans/import", files=files, headers={"Authorization": f"Bearer {admin_token}"})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["imported_exceptions"] == 1
    assert body["imported_ownership_rules"] == 1

    # The finding created in THIS SAME import must already reflect the exception and team.
    findings = client.get("/api/findings", headers={"Authorization": f"Bearer {admin_token}"}).json()["items"]
    assert len(findings) == 1
    finding = findings[0]
    assert finding["status"] == "exception"
    assert finding["team_name"] == "Postes de travail"

    # Exceptions and ownership rules are visible (and editable) via their own endpoints.
    exceptions = client.get("/api/exceptions", headers={"Authorization": f"Bearer {admin_token}"}).json()
    assert len(exceptions) == 1
    assert exceptions[0]["granted_by"] == "Eric Mercier"

    rules = client.get("/api/rules/ownership", headers={"Authorization": f"Bearer {admin_token}"}).json()
    assert len(rules) == 1
    assert rules[0]["hostname_pattern"] == "commence par q2"


def test_reimporting_does_not_duplicate_rules(client, admin_token):
    file_bytes = _build_workbook_with_rules()
    for _ in range(2):
        files = {
            "file": (
                "scan.xlsx",
                file_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        client.post("/api/scans/import", files=files, headers={"Authorization": f"Bearer {admin_token}"})

    exceptions = client.get("/api/exceptions", headers={"Authorization": f"Bearer {admin_token}"}).json()
    rules = client.get("/api/rules/ownership", headers={"Authorization": f"Bearer {admin_token}"}).json()
    assert len(exceptions) == 1
    assert len(rules) == 1
